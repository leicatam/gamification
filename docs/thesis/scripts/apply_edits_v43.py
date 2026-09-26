#!/usr/bin/env python3
import sys, json, copy, re, docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
src, dst, edits_f, report_f = sys.argv[1:5]
d = docx.Document(src); edits = json.load(open(edits_f)); report=[]
def norm(s): return s.replace('’',"'").replace('‘',"'").replace('“','"').replace('”','"').replace('–','-').replace('—','-').replace('\xa0',' ')
def find_variant(text, quote):
    if quote in text: return quote
    nt,nq=norm(text),norm(quote)
    if nq in nt and len(nt)==len(text):
        i=nt.index(nq); return text[i:i+len(quote)]
    return None
def replace_in_paragraph(p, quote, new):
    q=find_variant(p.text, quote)
    if q is None: return False,'quote not found'
    for r in p.runs:
        if q in r.text: r.text=r.text.replace(q,new,1); return True,'run-level'
    full=p.text.replace(q,new,1); runs=p.runs
    if not runs: return False,'no runs'
    runs[0].text=full
    for r in runs[1:]: r.text=''
    return True,'rebuilt'
def cell_paragraphs(loc):
    m=re.match(r'T(\d+)R(\d+)',loc); t=d.tables[int(m.group(1))]; row=t.rows[int(m.group(2))]
    seen=set(); ps=[]
    for c in row.cells:
        if id(c._tc) in seen: continue
        seen.add(id(c._tc)); ps.extend(c.paragraphs)
    return ps
# pre-resolve paragraph objects so insertions do not shift indices
paras=d.paragraphs
targets={}
for e in edits:
    if e['locator'].startswith('P'): targets[e['locator']]=paras[int(e['locator'][1:])]
def insert_after(p, text):
    anchor=p._element
    for chunk in text.split('\n'):
        new_el=copy.deepcopy(p._element); anchor.addnext(new_el); anchor=new_el
        np_=Paragraph(new_el,p._parent)
        for r in np_.runs[1:]: r._element.getparent().remove(r._element)
        # strip highlight/shading on the copied run
        if np_.runs:
            np_.runs[0].text=chunk
            rpr=np_.runs[0]._element.find(qn('w:rPr'))
            if rpr is not None:
                for tag in ('w:highlight','w:shd'):
                    x=rpr.find(qn(tag))
                    if x is not None: rpr.remove(x)
        else: np_.add_run(chunk)
        # remove paragraph-level shading if any
        ppr=new_el.find(qn('w:pPr'))
        if ppr is not None:
            x=ppr.find(qn('w:shd'))
            if x is not None: ppr.remove(x)
for e in edits:
    loc=e['locator']; act=e.get('action','reframe'); ok=False; how=''
    try:
        if loc.startswith('P'):
            p=targets[loc]
            if act=='reframe': ok,how=replace_in_paragraph(p,e['quote'],e['replacement'])
            elif act=='delete':
                q=find_variant(p.text,e['quote']) if e.get('quote') else None
                if q and q.strip()!=p.text.strip(): ok,how=replace_in_paragraph(p,e['quote'],'')
                else: p._element.getparent().remove(p._element); ok,how=True,'paragraph removed'
            elif act=='add_after': insert_after(p,e['replacement']); ok,how=True,'inserted after'
        elif loc.startswith('T'):
            for p in cell_paragraphs(loc):
                if find_variant(p.text,e['quote']):
                    ok,how=replace_in_paragraph(p,e['quote'],e['replacement'] if act!='delete' else ''); break
            if not ok: how='quote not found in row'
    except Exception as ex: ok,how=False,f'error: {ex}'
    report.append({'locator':loc,'src':e.get('src',''),'action':act,'ok':ok,'how':how,'quote':e['quote'][:70]})
d.save(dst); json.dump(report,open(report_f,'w'),indent=1,ensure_ascii=False)
bad=[r for r in report if not r['ok']]
print(f"applied {len(report)-len(bad)}/{len(report)}; failures {len(bad)}")
for r in bad: print('  FAIL',r['locator'],r['src'],r['how'],'|',r['quote'])

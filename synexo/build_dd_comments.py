#!/usr/bin/env python3
"""Comments memo on the FirstVitals Part-1 DD questionnaires
-> dd-comments.md + GenApep_DD_Questionnaire_Comments.docx.

To: Theresa Jang (WBI) and Brian Kim (EyeSel). Reviews the two identical
Part-1 questionnaires (v1.5, sections A-H) FirstVitals issued for its
proposed share-exchange/roll-up, and sets out how to respond without
accepting the frame. Grounded in the verified record (Plan v2) and a
3-agent readiness/red-team analysis of 17 Sep 2026.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DIR = "/home/user/gamification/synexo/"
NAVY = RGBColor(0x1B, 0x33, 0x55); GREY = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x2E, 0x86, 0xC1); RED = RGBColor(0xC0, 0x39, 0x2B)

CONF = ("PRIVATE & CONFIDENTIAL — for Theresa Jang (WBI) and Brian Kim (EyeSel). Not for "
        "release to FirstVitals or any third party. Read with GenApep IPO Plan v2.")
DISCLAIMER = ("Prepared 17 September 2026. Commercial comments for board discussion — not legal "
              "advice; response drafting and every 'None' answer should be settled with counsel.")

WHAT_SENT = [
    "FirstVitals has sent WBI and EyeSel two Part-1 'preliminary screening' DD questionnaires "
    "(v1.5, sections A–H: corporate, ownership, management, financials, valuation, products, "
    "regulatory, litigation). Part 2 — full documentary diligence — follows if the parties decide "
    "to proceed, 'toward the exchange ratio and definitive documentation.'",
    "The two documents are WORD-FOR-WORD IDENTICAL except the company name. That tells you three "
    "things: (a) it is boilerplate, not tailored diligence; (b) FV intends to DIFF your two "
    "responses against each other — treat drafting as one joint exercise, not two; (c) the level "
    "of tailoring is consistent with Plan v2's read of FV's sophistication.",
    "The questionnaire's own defined premise is the structure our documents rejected: 'FV would "
    "acquire the outstanding equity interests of the Company in exchange for shares of FV' — "
    "each Korean company rolled up SEPARATELY into FirstVitals as parent. Answering inside that "
    "defined term, without reservation, is soft acceptance of the frame.",
    "It is one-directional: FV asks you for three years of financials while our gate-zero "
    "questions to FV (1-A amounts actually sold, Form 1-K/1-SA runway, the $500K note's holder "
    "and terms, the FirstVitals/First Vital/ProteusDx entity-name chain, the 7.5M-share pool) "
    "remain unanswered — and NO NDA IS ON RECORD between the parties.",
]

POSITION = [
    ["Reserve the structure in the cover letter",
     "Respond jointly as one block. Re-define their term: 'the Proposed Transaction as described "
     "by FV, which the boards have not accepted or endorsed in structure, direction of "
     "acquisition, or valuation' — and use that term throughout, including H.3. Restate that any "
     "combination proceeds only through the JV holdco (block entry, pro rata, Addendum A "
     "machinery) and, per Plan v2, on the Management-Partnership basis: entity at verified "
     "value, Ernie's role and equity as disclosed KPI-gated management compensation."],
    ["NDA before anything sensitive moves",
     "Part 1's own instruction says formal supporting documentation is not required at this "
     "stage — use that: deliver concise figures-only answers now, defer every document (registry "
     "copies, statements, bank records, registers) to Part 2 under an executed mutual NDA and "
     "data-room protocol. Base the NDA condition on the boards' position, not on the "
     "third-party-consent instruction (reserve that for genuine third-party items: customer "
     "contracts, the counterparty MOU)."],
    ["Reciprocity as a condition of Part 2",
     "In a share-for-share exchange your shareholders become investors in FV stock — diligence "
     "is inherently mutual. Condition Part 2 on FV completing the IDENTICAL Part 1 about itself "
     "(gate-zero list above, plus management references for Mr. Lee including the 2024 "
     "shareholder dispute). No exchange-ratio discussion until both Part 1s are complete."],
    ["Cap the update duty",
     "Their instructions create a continuing-update expectation ('identify material changes "
     "after the date of the documents'). Signing Amendment No. 1, appointing the auditor, or "
     "closing the $2M bridge after delivery would all arguably qualify. State in the cover "
     "letter that responses speak as of a single stated date with no continuing duty outside an "
     "agreed Part-2 process — and sequence corporate actions versus delivery deliberately."],
]

RULES = [
    "One joint drafting protocol. Identical questionnaires get coordinated answers: same "
    "glossary, same as-of date, identical MOU-disclosure wording in the same places (E.3 with an "
    "A.3 cross-reference: 'a confidential non-binding MOU with a strategic partner regarding a "
    "potential business combination; terms subject to third-party consent'), and each company's "
    "written consent to the other's mention. Divergent wording reads as concealment and is "
    "exactly what the identical templates are designed to catch.",
    "Every 'None' gets a verification file. B.3, B.4, D.4, F.3, G.2, H.1, H.2: a dated internal "
    "memo naming who certified and what was searched (registry, court records, QA/CS, treasury, "
    "shareholder confirmations) BEFORE the answer is written. Use one uniform qualifier "
    "everywhere: 'None known to the Company after reasonable inquiry, as of [date].' These "
    "answers become the rep-and-warranty disclosure baseline of any later deal — and the "
    "defence file if there is none.",
    "Numbers: KRW primary, one FX rate, character-identical to the record. Answer in Korean "
    "(their instructions allow it) with Korean originals controlling; state figures in KRW with "
    "one disclosed rate consistent with the Amendment No. 1 FX-convention lock. Every figure "
    "must match the statutory/NTS record used for the 51:49 equalisation basis — a number given "
    "to FV that differs from the JV record is discoverable and weaponisable on both fronts.",
    "The $11.7M figure never appears. The DD report's '$11.7M 2026 WBI exosome revenue' is 3.4x "
    "the NTS-certified FY2025 actual and does not reconcile. If FV cites it — in Part 2, a term "
    "sheet, or a significance test — correct it in writing immediately (correct combined FY2025A "
    "is ~₩15,262.9M ≈ $10.9M @₩1,400).",
    "No internal documents, ever. Plan v2, the FirstVitals DD report, valuation scenarios "
    "($15–20M band, $8.0M peg, '~$50M combined'), route analysis, sponsor talks and dilution "
    "math stay out of every answer. Harmonised audit answer for both companies, verbatim: "
    "'Financial statements are unaudited, prepared under Korean SME accounting standards "
    "(중소기업회계기준); a PCAOB-registered audit engagement is being scoped for the Company's own "
    "purposes.' Nothing more.",
    "One response coordinator — not the conflicted channel. Route drafting through counsel or a "
    "neutral coordinator; the FA's documented conflicts (meeting advocacy, sweat-pool interest, "
    "FA mandate) make him the wrong sole channel for responses to the counterparty he introduced.",
]

WBI_TABLE = [
    ["A. Corporate", "READY / CAUTION",
     "Registry answers verbatim from a same-day 등기부 pull (all six A.2 documents still to be "
     "collected — defer copies to Part 2). A.3 perimeter: describe any excluded affiliate only "
     "from the registry, and verify the holdco/GenApep-Korea incorporation status as of the "
     "answer date; disclose the 12-Aug MOU as NON-BINDING and UNSIGNED-amendment status "
     "accurately (Amendment No. 1 is drafted, not signed)."],
    ["B. Ownership", "CAUTION",
     "Reconcile 111,448 shares / 5,572 placement / $89.75 against the 등기부 and 주주명부 BEFORE "
     "answering — these are user-supplied figures not yet verified against the registry. B.4: "
     "'None' at WBI level is accurate (the sweat pool is a holdco-level, unissued concept) — "
     "minute the reasoning."],
    ["C. Management", "NEEDS WORK",
     "Names and titles only, from the registry; no bios, no US-CEO/relocation discussion."],
    ["D. Financials", "CAUTION",
     "Only FY2025 figures are verified (rev ₩4,802.7M ≈ $3.43M; NI ₩243.2M). FY2023/FY2024 "
     "filings and H1-2026 YTD must be assembled — this is the same 'WBI evidence pack' Plan v2 "
     "already requires; build once, use for FV, Eyesel symmetry and the auditor. The implied "
     "₩172.3M operating profit stays OUT until read off the actual filing. D.5: frame the $2M "
     "bridge as growth/US-preparation funding by a modestly profitable company — neither "
     "distress nor full self-sufficiency."],
    ["E. Valuation", "CAUTION — the anchoring question",
     "E.1/E.2 answer: the 5% placement (5,572 sh at $89.75, ~$0.5M raised) — with an express "
     "non-reliance rider ('historic private placement; no valuation implication accepted for "
     "any exchange ratio'). Never the $15–20M band or the '$50M combined' remark. "
     "PRE-CONDITION: issue the one-pager corrective supplement to shareholders and complete the "
     "reliance check before the placement or the 주주명부 goes on paper."],
    ["F. Products", "CAUTION",
     "Stage precisely: exosome hair-loss line = commercialised, revenue-generating; MM Studio = "
     "in development, pre-revenue (the $2.90M 2026 figure is plan). Overstating MM Studio "
     "creates the same claims-vs-papers gap we condemned at FirstVitals. F.2: anonymised "
     "concentration bands now, names in Part 2 under NDA."],
    ["G. Regulatory", "CAUTION",
     "License register (MFDS registrations, GMP/quality, export certificates) with expiries; "
     "state the exosome products' regulatory CLASSIFICATION accurately (cosmetic vs "
     "functional/medical claims boundary). 5-year sweep incl. advertising-claims enforcement "
     "before any 'None'."],
    ["H. Litigation / catch-all", "CAUTION",
     "Court and internal checks covering officers/directors personally (first scope who counts "
     "as officer/director). H.3: flag narrowly — perimeter exclusion, JV framework and "
     "structure reservation; volunteer nothing about routes, valuations or FV's own condition."],
]

EYESEL_TABLE = [
    ["A. Corporate", "READY / NEEDS WORK",
     "Solvency answer is strong (profitable, dividend-paying, ₩10.36B retained earnings). All "
     "six A.2 documents still to be collected. A.3: K-SME standalone accounts suggest no "
     "consolidated entities — confirm before 'None'; the JV MOU has no incorporated entity yet."],
    ["B. Ownership", "CAUTION",
     "Build the cap table from the 주주명부 and explain the ₩478,000 share-premium transfer "
     "consistently with the registry. B.2/C.1 are THE authority items: the registry states the "
     "representative director of record; obtain the written board confirmation of the principal "
     "(Park vs Kim) BEFORE any answer or signature leaves — the whole response is tainted if "
     "signed by the wrong person. B.4 'None' (after 정관/minutes check) sharpens the demand that "
     "FV's 54.18%-FD pool be cancelled."],
    ["C. Management", "CAUTION",
     "Same authority point; headcount from HR with a stated date (SG&A salaries ₩669.8M give "
     "the sanity anchor; manufacturing labour sits in COGS)."],
    ["D. Financials", "STRONGEST HAND / GAPS",
     "FY2025/FY2024 income statement and retained-earnings statement are verified and strong. "
     "Missing: balance sheet, FY2023, 2026 YTD, cash-flow statement (may legitimately not exist "
     "under K-SME — say so, don't fabricate one), CIT return, AR/AP aging, debt schedule, bank "
     "statements. Staple the OP-decline bridge to D.2 — INCLUDING the gross-margin compression "
     "(-6.1pp, ≈₩642M at FY2025 volume), not just SG&A growth — and pre-normalise NI for the "
     "₩301.2M one-off disposal gain before FV does it adversarially. ₩10.36B retained earnings "
     "is NOT cash (interest income halved; depreciation tripled) — no cash figure until bank "
     "statements are in hand."],
    ["E. Valuation", "CAUTION",
     "E.1: 'no independent third-party valuation of the Company exists' — exactly that, minuted "
     "with the reasoning (internal planning bands are negotiating positions, not valuations). "
     "E.2: base 'no financings' on the 등기부 issuance history 2023–2026, not on the income "
     "statement's silence. E.3: the joint MOU-disclosure wording, mirrored with WBI."],
    ["F. Products", "NEEDS WORK / CAUTION",
     "Everything is launched and revenue-generating — a strength against a pre-revenue "
     "counterparty; draft the one-paragraph description aligned to the licence register. Verify "
     "actual GMP certifications (scope, number, expiry) before the narrative uses 'GMP'. F.2 "
     "top-10 customers: anonymised bands now, names in Part 2 under NDA; check whether the "
     "₩435.0M commission counterparty appears in or near that list."],
    ["G. Regulatory", "NEEDS WORK",
     "Licence register incl. export registrations covering the ₩5,294.0M export book; check "
     "change-of-control transferability of each licence (a Part-2 question to know first); "
     "grant-clawback conditions on the ₩62.7M government subsidies; the 4.1% ETR explanation "
     "(SME/R&D credits) ready before the CIT return is delivered."],
    ["H. Litigation / catch-all", "CAUTION",
     "H.2 CANNOT be answered until the ₩435.0M FY2025 sales-commission counterparty (10.5x YoY) "
     "is understood internally — related-party, undocumented intermediary or rebate arrangement "
     "would land exactly here, and matters in any future US-listing context. Resolve before the "
     "response is sent. H.3: narrow, accurate, with the structure reservation."],
]

INSTRUCTION_TRAPS = [
    "'Three most recent fiscal years' = FY2023–FY2025 plus YTD: neither company has FY2023 "
    "material assembled. Inventory what exists NOW and pre-draft the unavailability "
    "explanations their instructions require, rather than promising three years and failing.",
    "The consolidation sweep: 'Company' includes subsidiaries, controlled affiliates and "
    "consolidated entities — so every 'None' is deemed given for affiliates too. Preface: each "
    "company answers on a standalone K-SME basis, no consolidated entities as of the answer "
    "date. Note that a later-incorporated holdco/GenApep Korea would fall INSIDE their "
    "definition — another reason to cap the update duty.",
    "Part 1 says documents are not required, yet A.2/D.1 request copies 'for FV's CFO'. Invoke "
    "the conflict explicitly: figures-only Part 1, documents in Part 2 under NDA.",
    "B.2 (ultimate beneficial owner): state the standard applied (e.g., Korean AML 25% "
    "ownership/control test) inside the answer, so it cannot be re-read later against a "
    "different definition.",
    "Estimates: Part 1 invites 'good-faith estimates' — a trap for D.3/D.5-type items that later "
    "become rep baselines. Estimate only what is safe to estimate (headcount), never balance-"
    "sheet or debt figures.",
]

PART2_PREP = [
    "Their Part 1 omits what a competent Part 2 will demand — prepare now: IP/chain-of-title "
    "(our own Amendment No. 1 condition precedent anyway: AI.pep, CodeLife.AI, IT-EXO, SynExo "
    "assignments), related-party transactions (both companies — and note their questionnaire's "
    "own silence on related parties), employment/labor, data privacy, insurance schedules.",
    "EyeSel: resolve the ₩435.0M commission counterparty and the severance-liability position "
    "(the missing balance sheet's classic K-SME item) before Part 2 exposes them.",
    "WBI: MM Studio contract evidence and H1-2026 actuals — the same pack the 49% defence and "
    "the auditor need.",
    "Both: the PCAOB audit workstream (Q4-2026 appointment) produces most Part-2 financial "
    "artefacts as a by-product — sequence the data room off the audit file, not as a separate "
    "scramble.",
]

SEQUENCE = [
    ["Week 1", "Joint board minute adopting the response position (structure reservation, NDA-first, reciprocity, update-duty cap); mutual NDA executed; written confirmation of the EyeSel principal; WBI corrective supplement to shareholders issued."],
    ["Weeks 1–2", "Assemble: registry extracts, cap-table reconciliations, FY2023 inventory, WBI evidence pack, EyeSel balance sheet + bank statements; verification files for every intended 'None'; resolve the commission counterparty internally."],
    ["Weeks 2–3", "Deliver coordinated figures-only Part-1 responses under the joint cover letter; simultaneously deliver our reciprocal Part-1 to FV (gate-zero list); documents deferred to Part 2 under NDA/data-room protocol."],
    ["Gate", "No exchange-ratio or valuation discussion until FV's reciprocal Part 1 is complete and gate-zero closes; any Part-2 process runs off the audit file."],
]

# ------------------------------------------------------------------ markdown
def wmd():
    L = []; a = L.append
    a("# Comments for Theresa & Brian Kim — the FirstVitals Part-1 DD questionnaires\n")
    a("> **%s**\n>\n> %s\n" % (CONF, DISCLAIMER))
    a("\n## 1. What FirstVitals sent — and what it signals\n")
    for t in WHAT_SENT: a("- %s\n" % t)
    a("\n## 2. The response position (answer without accepting the frame)\n")
    for t, d in POSITION: a("- **%s** — %s\n" % (t, d))
    a("\n## 3. Six rules before anything leaves the building\n")
    for i, r in enumerate(RULES, 1): a("%d. %s\n" % (i, r))
    a("\n## 4. WBI answer-readiness (for Theresa)\n\n| Section | Status | Key actions & traps |\n|---|---|---|\n")
    for r in WBI_TABLE: a("| **%s** | %s | %s |\n" % (r[0], r[1], r[2]))
    a("\n## 5. EyeSel answer-readiness (for Brian Kim)\n\n| Section | Status | Key actions & traps |\n|---|---|---|\n")
    for r in EYESEL_TABLE: a("| **%s** | %s | %s |\n" % (r[0], r[1], r[2]))
    a("\n## 6. Traps inside the questionnaire's own instructions\n")
    for t in INSTRUCTION_TRAPS: a("- %s\n" % t)
    a("\n## 7. What Part 2 will bring — prepare now\n")
    for t in PART2_PREP: a("- %s\n" % t)
    a("\n## 8. Suggested sequence\n\n| When | Actions |\n|---|---|\n")
    for r in SEQUENCE: a("| **%s** | %s |\n" % (r[0], r[1]))
    a("\n---\n*Basis: 3-agent readiness analysis with adversarial red-team over the questionnaires "
      "and the verified record (17 Sep 2026). Answering Part 1 on these terms commits the "
      "companies to nothing and runs in parallel with Stage 1.*\n")
    open(DIR + "dd-comments.md", "w").write("".join(L))

wmd()

# ------------------------------------------------------------------ docx
def nd():
    doc = Document()
    n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(10.5)
    for lvl, sz in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        st = doc.styles[lvl]; st.font.name = "Calibri"; st.font.size = Pt(sz)
        st.font.color.rgb = NAVY; st.font.bold = True
    return doc

def para(doc, t="", size=10.5, bold=False, italic=False, color=None, align=None, after=6):
    p = doc.add_paragraph()
    if align: p.alignment = align
    r = p.add_run(t); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after); return p

def bullet(doc, t, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if lead:
        r = p.add_run(lead); r.bold = True
    p.add_run(t); p.paragraph_format.space_after = Pt(3); return p

def numbered(doc, t):
    p = doc.add_paragraph(style="List Number")
    p.add_run(t); p.paragraph_format.space_after = Pt(3); return p

def tbl(doc, headers, rows, widths=None, fs=9):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        r = t.rows[0].cells[i].paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(fs)
    for row in rows:
        c = t.add_row().cells
        for i, v in enumerate(row):
            rr = c[i].paragraphs[0].add_run(str(v)); rr.font.size = Pt(fs)
            if i == 0: rr.bold = True
    if widths:
        for i, w in enumerate(widths):
            for r in t.rows: r.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t

doc = nd()
para(doc, "GENAPEP", size=26, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(doc, "Comments on the FirstVitals DD Questionnaires (Part 1) — for Theresa Jang & Brian Kim",
     size=14, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(doc, CONF, size=9, bold=True, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para(doc, DISCLAIMER, size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

doc.add_heading("1. What FirstVitals sent — and what it signals", level=2)
for t in WHAT_SENT:
    bullet(doc, t)

doc.add_heading("2. The response position (answer without accepting the frame)", level=2)
for t, d in POSITION:
    bullet(doc, d, t + " — ")

doc.add_heading("3. Six rules before anything leaves the building", level=2)
for r in RULES:
    numbered(doc, r)

doc.add_heading("4. WBI answer-readiness (for Theresa)", level=2)
tbl(doc, ["Section", "Status", "Key actions & traps"], WBI_TABLE, widths=[1.2, 1.0, 4.3])

doc.add_heading("5. EyeSel answer-readiness (for Brian Kim)", level=2)
tbl(doc, ["Section", "Status", "Key actions & traps"], EYESEL_TABLE, widths=[1.2, 1.0, 4.3])

doc.add_heading("6. Traps inside the questionnaire's own instructions", level=2)
for t in INSTRUCTION_TRAPS:
    bullet(doc, t)

doc.add_heading("7. What Part 2 will bring — prepare now", level=2)
for t in PART2_PREP:
    bullet(doc, t)

doc.add_heading("8. Suggested sequence", level=2)
tbl(doc, ["When", "Actions"], SEQUENCE, widths=[1.1, 5.4])
para(doc, "Basis: 3-agent readiness analysis with adversarial red-team over the questionnaires "
          "and the verified record (17 September 2026). Answering Part 1 on these terms commits "
          "the companies to nothing and runs in parallel with Stage 1.",
     size=8.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(DIR + "GenApep_DD_Questionnaire_Comments.docx")
print("Saved dd-comments.md, GenApep_DD_Questionnaire_Comments.docx")
